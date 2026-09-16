"""
src/ingestion/loader.py
========================
Stage 1 of the pipeline: INGESTION + PARSING.

Turns an arbitrary uploaded document (.pdf, .docx, .txt) into a list of
`RawElement` objects - each one tagged with its TYPE (text / table / image)
and its source location (page number, etc). This element-level parsing is
what makes multimodality possible: text, tables, and images are kept
separate and processed with the strategy that suits each one, instead of
being flattened into one lossy text blob (which is what naive
`pdf_to_text()` approaches do, and why they fail on scanned pages, tables,
and diagrams).

Supported inputs:
  - PDF (digital-native, scanned/image-only, or mixed) via PyMuPDF + pdfplumber
  - DOCX via python-docx
  - Plain text / markdown

Design note: this is deliberately dependency-light (no `unstructured`
library) so the project stays easy to install and deploy on a free tier.
"""

from __future__ import annotations
import io
import os
from dataclasses import dataclass, field
from typing import Optional

import fitz  # PyMuPDF
import pdfplumber
from docx import Document as DocxDocument
from PIL import Image

from utils.logger import get_logger

logger = get_logger("ingestion.loader")


@dataclass
class RawElement:
    """One atomic piece of content pulled out of a document."""
    element_type: str          # "text" | "table" | "image" | "scanned_page"
    content: str                # text content (markdown table for tables;
                                 # empty for images until captioned)
    page_number: int
    source_file: str
    image: Optional[Image.Image] = None   # populated only for image/scanned_page
    extra: dict = field(default_factory=dict)


class DocumentLoader:
    """Routes a file to the correct parser based on extension."""

    def load(self, file_path: str) -> list[RawElement]:
        ext = os.path.splitext(file_path)[1].lower()
        filename = os.path.basename(file_path)

        if ext == ".pdf":
            return self._load_pdf(file_path, filename)
        elif ext == ".docx":
            return self._load_docx(file_path, filename)
        elif ext in (".txt", ".md"):
            return self._load_text(file_path, filename)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

    # ------------------------------------------------------------------
    # PDF - the most complex case, because a single PDF can mix real text,
    # embedded raster images, vector-drawn tables, and fully-scanned pages
    # (no text layer at all).
    # ------------------------------------------------------------------
    def _load_pdf(self, file_path: str, filename: str) -> list[RawElement]:
        elements: list[RawElement] = []
        doc = fitz.open(file_path)

        # Pass 1: text (via PyMuPDF) + embedded images, page by page.
        for page_index in range(len(doc)):
            page = doc[page_index]
            page_num = page_index + 1
            text = page.get_text("text").strip()

            if text:
                elements.append(
                    RawElement(
                        element_type="text",
                        content=text,
                        page_number=page_num,
                        source_file=filename,
                    )
                )
            else:
                # No extractable text layer -> this page is very likely a
                # scanned image. Render the whole page as an image so the
                # OCR / vision-captioning stage can handle it.
                pix = page.get_pixmap(dpi=200)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                elements.append(
                    RawElement(
                        element_type="scanned_page",
                        content="",
                        page_number=page_num,
                        source_file=filename,
                        image=img,
                    )
                )
                logger.info(f"Page {page_num} has no text layer -> flagged as scanned_page.")

            # Embedded raster images (diagrams, logos, photos pasted into
            # the PDF, coverage flowcharts, etc.)
            for img_index, img_info in enumerate(page.get_images(full=True)):
                xref = img_info[0]
                try:
                    base_image = doc.extract_image(xref)
                    img_bytes = base_image["image"]
                    pil_img = Image.open(io.BytesIO(img_bytes))
                    # Skip tiny images - these are almost always icons/
                    # bullet graphics/logos with no informational content,
                    # and captioning them just adds noise to the index.
                    if pil_img.width < 80 or pil_img.height < 80:
                        continue
                    elements.append(
                        RawElement(
                            element_type="image",
                            content="",
                            page_number=page_num,
                            source_file=filename,
                            image=pil_img,
                            extra={"image_index": img_index},
                        )
                    )
                except Exception as e:
                    logger.warning(f"Could not extract image {img_index} on page {page_num}: {e}")

        doc.close()

        # Pass 2: tables (pdfplumber is noticeably better than PyMuPDF at
        # detecting ruled/bordered tables and reconstructing row/column
        # structure).
        try:
            with pdfplumber.open(file_path) as pdf:
                for page_index, page in enumerate(pdf.pages):
                    page_num = page_index + 1
                    tables = page.extract_tables()
                    for t_index, table in enumerate(tables):
                        md_table = self._table_to_markdown(table)
                        if md_table:
                            elements.append(
                                RawElement(
                                    element_type="table",
                                    content=md_table,
                                    page_number=page_num,
                                    source_file=filename,
                                    extra={"table_index": t_index},
                                )
                            )
        except Exception as e:
            logger.warning(f"Table extraction failed: {e}")

        logger.info(
            f"Parsed '{filename}': "
            f"{sum(1 for e in elements if e.element_type == 'text')} text blocks, "
            f"{sum(1 for e in elements if e.element_type == 'table')} tables, "
            f"{sum(1 for e in elements if e.element_type == 'image')} images, "
            f"{sum(1 for e in elements if e.element_type == 'scanned_page')} scanned pages."
        )
        return elements

    @staticmethod
    def _table_to_markdown(table: list[list]) -> str:
        """Convert pdfplumber's raw table (list of rows) into a markdown
        table string. Markdown tables embed cleanly into LLM prompts and
        preserve row/column relationships far better than flattened text."""
        if not table or len(table) < 1:
            return ""
        rows = [[(cell or "").strip().replace("\n", " ") for cell in row] for row in table]
        header, *body = rows
        if not any(header):
            return ""
        md = "| " + " | ".join(header) + " |\n"
        md += "| " + " | ".join(["---"] * len(header)) + " |\n"
        for row in body:
            row = row + [""] * (len(header) - len(row))  # pad ragged rows
            md += "| " + " | ".join(row[: len(header)]) + " |\n"
        return md

    # ------------------------------------------------------------------
    # DOCX
    # ------------------------------------------------------------------
    def _load_docx(self, file_path: str, filename: str) -> list[RawElement]:
        elements: list[RawElement] = []
        doc = DocxDocument(file_path)

        # python-docx has no real "page" concept, so we use paragraph-group
        # index as a stand-in locator for citation purposes.
        buffer = []
        block_num = 1
        for para in doc.paragraphs:
            if para.text.strip():
                buffer.append(para.text.strip())
            if len("\n".join(buffer)) > 1200:  # flush into a block
                elements.append(
                    RawElement("text", "\n".join(buffer), block_num, filename)
                )
                buffer = []
                block_num += 1
        if buffer:
            elements.append(RawElement("text", "\n".join(buffer), block_num, filename))

        # Tables in docx
        for t_index, table in enumerate(doc.tables):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            md_table = self._table_to_markdown(rows)
            if md_table:
                elements.append(
                    RawElement("table", md_table, block_num, filename, extra={"table_index": t_index})
                )

        logger.info(f"Parsed '{filename}': {len(elements)} elements from DOCX.")
        return elements

    # ------------------------------------------------------------------
    # Plain text / markdown
    # ------------------------------------------------------------------
    def _load_text(self, file_path: str, filename: str) -> list[RawElement]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        return [RawElement("text", text, 1, filename)]
